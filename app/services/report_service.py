from __future__ import annotations

from datetime import datetime
from pathlib import Path
from uuid import UUID, uuid4

from docx import Document
from docx.shared import Inches, Pt

from app.core.errors import ReportGenerationError
from app.core.policies import inside
from app.core.schemas import DatasetSummary


class ReportService:
    def generate(self, output_root: Path, session_id: UUID, file_name: str, summary: DatasetSummary,
                 narrative: str, chart: Path | None = None) -> Path:
        session_dir = inside(output_root, output_root / str(session_id))
        session_dir.mkdir(parents=True, exist_ok=True)
        target = inside(session_dir, session_dir / f"report-{uuid4()}.docx")
        try:
            document = Document()
            normal = document.styles["Normal"]
            normal.font.name = "Arial"; normal.font.size = Pt(10.5)
            title = document.add_heading("Отчёт по анализу данных", 0)
            title.style = document.styles["Title"]
            document.add_paragraph(f"Файл: {file_name}")
            document.add_paragraph(f"Дата формирования: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")
            document.add_heading("Краткий вывод", level=1)
            document.add_paragraph(narrative)
            document.add_heading("Ключевые метрики", level=1)
            table = document.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "Показатель"; table.rows[0].cells[1].text = "Значение"
            for label, value in (("Строки", summary.row_count), ("Столбцы", summary.column_count)):
                row = table.add_row().cells; row[0].text = label; row[1].text = str(value)
            document.add_heading("Качество данных", level=1)
            notes = [f"{column['name']}: пропусков {column['missing']}, уникальных {column['unique']}" for column in summary.columns[:12]]
            document.add_paragraph("; ".join(notes) or "Поля не определены.")
            document.add_heading("Preview", level=1)
            if summary.preview:
                headers = list(summary.preview[0].keys())[:8]
                preview_table = document.add_table(rows=1, cols=len(headers)); preview_table.style = "Table Grid"
                for index, header in enumerate(headers): preview_table.rows[0].cells[index].text = str(header)
                for item in summary.preview[:5]:
                    cells = preview_table.add_row().cells
                    for index, header in enumerate(headers): cells[index].text = str(item.get(header, ""))[:80]
            if chart and chart.exists():
                document.add_heading("График", level=1)
                document.add_picture(str(chart), width=Inches(6.2))
            document.add_heading("Summary", level=1)
            document.add_paragraph("Отчёт содержит агрегированные показатели и ограниченный preview, без полной выгрузки исходного набора данных.")
            document.save(target)
        except Exception as exc:
            target.unlink(missing_ok=True)
            raise ReportGenerationError() from exc
        return target
